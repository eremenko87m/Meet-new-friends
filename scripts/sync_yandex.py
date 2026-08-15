from __future__ import annotations

import io
import json
import mimetypes
import posixpath
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile

import requests
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple


ROOT = Path(__file__).resolve().parents[1]

CONFIG = json.loads(
    (ROOT / "config.json").read_text(
        encoding="utf-8"
    )
)

DATA_PATH = ROOT / "data" / "exercises.json"
MEDIA_ROOT = ROOT / "assets" / "media"

YANDEX_META_API = (
    "https://cloud-api.yandex.net/"
    "v1/disk/public/resources"
)

YANDEX_DOWNLOAD_API = (
    "https://cloud-api.yandex.net/"
    "v1/disk/public/resources/download"
)

TIMEOUT = 45

NO_CACHE_HEADERS = {
    "Cache-Control": "no-cache, no-store, max-age=0",
    "Pragma": "no-cache",
    "User-Agent": "ExerciseLibrarySync/4.0",
}


# =========================================================
# YANDEX DOWNLOAD
# =========================================================


def yandex_metadata(public_url: str) -> dict:

    r = requests.get(
        YANDEX_META_API,
        params={
            "public_key": public_url,
            "fields": (
                "name,modified,size,"
                "type,mime_type"
            ),
        },
        headers=NO_CACHE_HEADERS,
        timeout=TIMEOUT,
    )

    r.raise_for_status()

    return r.json()


def yandex_direct(public_url: str) -> str:

    r = requests.get(
        YANDEX_DOWNLOAD_API,
        params={
            "public_key": public_url
        },
        headers=NO_CACHE_HEADERS,
        timeout=TIMEOUT,
    )

    r.raise_for_status()

    href = r.json().get("href")

    if not href:

        raise RuntimeError(
            "Yandex did not return "
            "a download URL for "
            f"{public_url}"
        )

    return href


def download_public_file(
    public_url: str
) -> bytes:

    meta = yandex_metadata(
        public_url
    )

    print(
        "Yandex source:",
        f"name={meta.get('name')!r}",
        f"modified={meta.get('modified')!r}",
        f"size={meta.get('size')!r}",
        f"type={meta.get('type')!r}",
    )

    href = yandex_direct(
        public_url
    )

    r = requests.get(
        href,
        headers=NO_CACHE_HEADERS,
        timeout=TIMEOUT,
    )

    r.raise_for_status()

    content = r.content

    if len(content) < 100:

        raise RuntimeError(
            "Downloaded Yandex file "
            "is unexpectedly small: "
            f"{len(content)} bytes"
        )

    # XLSX is actually a ZIP archive,
    # so it normally starts with PK.

    if not content.startswith(b"PK"):

        raise RuntimeError(
            "The public Yandex link "
            "did not download an XLSX file. "
            "Content-Type="
            f"{r.headers.get('content-type')!r}"
        )

    print(
        "Downloaded workbook:",
        len(content),
        "bytes"
    )

    return content


# =========================================================
# FIX YANDEX XLSX
# =========================================================


def sanitize_broken_table_references(
    content: bytes
) -> bytes:
    """
    Yandex can export an XLSX with
    references to Excel tables such as:

    xl/tables/table1.xml

    even though the actual table1.xml
    file is missing.

    openpyxl then crashes with:

    KeyError:
    There is no item named
    'xl/tables/table1.xml'

    We remove ONLY those broken
    table references.

    Cell values and pictures remain.
    """

    REL_NS = (
        "http://schemas.openxmlformats.org/"
        "package/2006/relationships"
    )

    MAIN_NS = (
        "http://schemas.openxmlformats.org/"
        "spreadsheetml/2006/main"
    )

    DOC_REL_NS = (
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships"
    )

    ET.register_namespace(
        "",
        REL_NS
    )

    ET.register_namespace(
        "r",
        DOC_REL_NS
    )

    source = io.BytesIO(
        content
    )

    with ZipFile(
        source,
        "r"
    ) as zin:

        names = set(
            zin.namelist()
        )

        replacements = {}

        removed_total = 0

        for rels_name in names:

            if not (
                rels_name.startswith(
                    "xl/worksheets/_rels/"
                )
                and rels_name.endswith(
                    ".xml.rels"
                )
            ):
                continue

            rel_root = ET.fromstring(
                zin.read(
                    rels_name
                )
            )

            removed_ids = set()

            changed = False

            # Example:
            #
            # sheet1.xml.rels
            # ->
            # sheet1.xml

            sheet_filename = (
                rels_name
                .rsplit("/", 1)[-1][:-5]
            )

            sheet_name = (
                "xl/worksheets/"
                + sheet_filename
            )

            for rel in list(
                rel_root
            ):

                rel_type = (
                    rel.attrib.get(
                        "Type",
                        ""
                    )
                )

                if not rel_type.endswith(
                    "/table"
                ):
                    continue

                target = (
                    rel.attrib.get(
                        "Target",
                        ""
                    )
                )

                if target.startswith(
                    "/"
                ):

                    target_path = (
                        target.lstrip("/")
                    )

                else:

                    target_path = (
                        posixpath.normpath(
                            posixpath.join(
                                posixpath.dirname(
                                    sheet_name
                                ),
                                target,
                            )
                        )
                    )

                # If target exists,
                # everything is fine.

                if target_path in names:
                    continue

                relationship_id = (
                    rel.attrib.get(
                        "Id",
                        ""
                    )
                )

                print(
                    "Removing broken XLSX "
                    "table relationship:",
                    rels_name,
                    relationship_id,
                    "->",
                    target_path,
                )

                removed_ids.add(
                    relationship_id
                )

                rel_root.remove(
                    rel
                )

                changed = True

                removed_total += 1

            if not changed:
                continue

            replacements[
                rels_name
            ] = ET.tostring(
                rel_root,
                encoding="utf-8",
                xml_declaration=True,
            )

            # Remove corresponding
            # <tablePart> from worksheet.

            if (
                sheet_name in names
                and removed_ids
            ):

                sheet_root = (
                    ET.fromstring(
                        zin.read(
                            sheet_name
                        )
                    )
                )

                table_parts = (
                    sheet_root.find(
                        f"{{{MAIN_NS}}}"
                        "tableParts"
                    )
                )

                if (
                    table_parts
                    is not None
                ):

                    for table_part in list(
                        table_parts
                    ):

                        rid = (
                            table_part
                            .attrib
                            .get(
                                f"{{{DOC_REL_NS}}}"
                                "id"
                            )
                        )

                        if rid in removed_ids:

                            table_parts.remove(
                                table_part
                            )

                    remaining = list(
                        table_parts
                    )

                    if not remaining:

                        sheet_root.remove(
                            table_parts
                        )

                    else:

                        table_parts.set(
                            "count",
                            str(
                                len(
                                    remaining
                                )
                            ),
                        )

                    replacements[
                        sheet_name
                    ] = ET.tostring(
                        sheet_root,
                        encoding="utf-8",
                        xml_declaration=True,
                    )

        if removed_total == 0:

            print(
                "No broken XLSX table "
                "relationships found."
            )

            return content

        print(
            "Removed",
            removed_total,
            "broken XLSX table "
            "relationship(s)."
        )

        output = io.BytesIO()

        with ZipFile(
            output,
            "w",
            ZIP_DEFLATED
        ) as zout:

            for info in (
                zin.infolist()
            ):

                data = (
                    replacements.get(
                        info.filename,
                        zin.read(
                            info.filename
                        ),
                    )
                )

                zout.writestr(
                    info,
                    data
                )

        return output.getvalue()


# =========================================================
# BASIC HELPERS
# =========================================================


def clean_text(value) -> str:

    if value is None:
        return ""

    if isinstance(
        value,
        bool
    ):

        return (
            "YES"
            if value
            else "NO"
        )

    return str(
        value
    ).strip()


def is_yandex_share(
    url: str
) -> bool:

    try:

        host = (
            urlparse(url)
            .netloc
            .lower()
        )

        return (
            host.endswith(
                "disk.yandex.ru"
            )
            or host.endswith(
                "yadi.sk"
            )
        )

    except Exception:

        return False


def ext_from_response(
    resp: requests.Response,
    fallback: str,
) -> str:

    cd = resp.headers.get(
        "content-disposition",
        "",
    )

    m = re.search(
        r"filename\*?="
        r"(?:UTF-8'')?"
        r"[\"']?"
        r"([^\"';]+)",
        cd,
        re.I,
    )

    if m:

        suffix = Path(
            m.group(1)
        ).suffix

        if suffix:
            return suffix.lower()

    ct = (
        resp.headers
        .get(
            "content-type",
            ""
        )
        .split(";")[0]
        .strip()
    )

    return (
        mimetypes.guess_extension(
            ct
        )
        or fallback
    )


# =========================================================
# MEDIA FROM LINKS
# =========================================================


def clear_old_media(
    section: str,
    no: int,
    kind: str,
) -> None:

    folder = (
        MEDIA_ROOT
        / section
    )

    if not folder.exists():
        return

    for old in folder.glob(
        f"{no:02d}-{kind}.*"
    ):

        old.unlink(
            missing_ok=True
        )


def cache_media(
    url: str,
    section: str,
    no: int,
    kind: str,
) -> str:

    url = clean_text(
        url
    )

    if not url:
        return ""

    # Normal external URL:
    # leave as-is.

    if not is_yandex_share(
        url
    ):

        return url

    # Yandex public link:
    # download media into site.

    try:

        href = yandex_direct(
            url
        )

        resp = requests.get(
            href,
            headers=NO_CACHE_HEADERS,
            timeout=TIMEOUT,
        )

        resp.raise_for_status()

        fallback = (
            ".jpg"
            if kind == "image"
            else ".mp3"
        )

        ext = ext_from_response(
            resp,
            fallback,
        )

        folder = (
            MEDIA_ROOT
            / section
        )

        folder.mkdir(
            parents=True,
            exist_ok=True,
        )

        clear_old_media(
            section,
            no,
            kind,
        )

        path = (
            folder
            / f"{no:02d}-{kind}{ext}"
        )

        path.write_bytes(
            resp.content
        )

        return (
            path
            .relative_to(ROOT)
            .as_posix()
        )

    except Exception as exc:

        print(
            "WARNING media",
            url,
            ":",
            exc
        )

        return url


# =========================================================
# EMBEDDED / FLOATING PICTURES FROM EXCEL
# =========================================================


def image_anchor_position(
    img
):
    """
    Returns:

    (excel_row, excel_col)

    Both are 1-based.
    """

    anchor = getattr(
        img,
        "anchor",
        None,
    )

    if anchor is None:
        return None

    # Sometimes anchor can
    # simply be "D2".

    if isinstance(
        anchor,
        str
    ):

        try:

            return (
                coordinate_to_tuple(
                    anchor
                )
            )

        except Exception:

            return None

    # Normal openpyxl anchor.

    marker = getattr(
        anchor,
        "_from",
        None,
    )

    if marker is not None:

        try:

            return (
                int(
                    marker.row
                ) + 1,
                int(
                    marker.col
                ) + 1,
            )

        except Exception:

            return None

    return None


def embedded_image_bytes(
    img
):

    data_method = getattr(
        img,
        "_data",
        None,
    )

    if not callable(
        data_method
    ):

        raise RuntimeError(
            "openpyxl image object "
            "has no readable image data"
        )

    raw = data_method()

    if not raw:

        raise RuntimeError(
            "embedded image "
            "has no data"
        )

    fmt = clean_text(
        getattr(
            img,
            "format",
            "",
        )
    ).lower().lstrip(".")

    if fmt == "jpg":
        fmt = "jpeg"

    if fmt not in {
        "png",
        "jpeg",
        "gif",
        "bmp",
        "tiff",
    }:

        fmt = "png"

    ext = (
        ".jpg"
        if fmt == "jpeg"
        else f".{fmt}"
    )

    return raw, ext


def embedded_images_for_rows(
    ws,
    image_col_index: int | None,
    cards_per: int,
):
    """
    Find floating pictures.

    Picture is assigned to card
    according to the row where
    its top-left corner sits.

    Examples:

    D2 = card 1
    D3 = card 2
    D4 = card 3
    """

    images = list(
        getattr(
            ws,
            "_images",
            [],
        )
        or []
    )

    if not images:

        print(
            f"{ws.title}: "
            "0 embedded picture(s)"
        )

        return {}

    candidates = defaultdict(
        list
    )

    for order, img in enumerate(
        images
    ):

        pos = image_anchor_position(
            img
        )

        if not pos:

            print(
                f"WARNING {ws.title}: "
                "embedded picture "
                f"#{order + 1} "
                "has no readable anchor"
            )

            continue

        excel_row, excel_col = pos

        # Row 1 = header.
        # Exercise rows begin at 2.

        if (
            2
            <= excel_row
            <= cards_per + 1
        ):

            # row_map uses zero-based
            # indexes, Excel columns
            # are one-based.

            target_col = (
                image_col_index + 1
                if image_col_index
                is not None
                else excel_col
            )

            distance = abs(
                excel_col
                - target_col
            )

            candidates[
                excel_row
            ].append(
                (
                    distance,
                    order,
                    excel_col,
                    img,
                )
            )

    chosen = {}

    for (
        excel_row,
        row_candidates,
    ) in candidates.items():

        # If several pictures exist
        # in same row, use picture
        # closest to Image column.

        row_candidates.sort(
            key=lambda item: (
                item[0],
                item[1],
            )
        )

        (
            distance,
            _,
            excel_col,
            img,
        ) = row_candidates[0]

        chosen[
            excel_row
        ] = img

        if (
            image_col_index
            is not None
            and distance
        ):

            print(
                f"NOTE {ws.title}: "
                f"picture in row "
                f"{excel_row} "
                "is anchored in "
                f"column {excel_col}; "
                "using it for Image "
                f"column "
                f"{image_col_index + 1}."
            )

    print(
        f"{ws.title}: "
        f"{len(images)} "
        "embedded picture(s), "
        f"{len(chosen)} "
        "mapped to exercise row(s)"
    )

    return chosen


def save_embedded_image(
    img,
    section: str,
    no: int,
) -> str:

    raw, ext = (
        embedded_image_bytes(
            img
        )
    )

    folder = (
        MEDIA_ROOT
        / section
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    clear_old_media(
        section,
        no,
        "image",
    )

    path = (
        folder
        / f"{no:02d}-image{ext}"
    )

    path.write_bytes(
        raw
    )

    return (
        path
        .relative_to(ROOT)
        .as_posix()
    )


# =========================================================
# TABLE HEADERS
# =========================================================


def row_map(
    headers
):

    norm = {
        clean_text(v).lower(): i
        for i, v
        in enumerate(headers)
    }

    aliases = {

        "no": [
            "no",
            "#",
            "number",
            "номер",
            "№",
        ],

        "title": [
            "title",
            "name",
            "название",
        ],

        "link": [
            "link",
            "url",
            "ссылка",
        ],

        "image": [
            "image",
            "picture",
            "img",
            "картинка",
        ],

        "audio": [
            "audio",
            "аудио",
        ],

        "level": [
            "level",
            "уровень",
        ],

        "notes": [
            "notes",
            "note",
            "comment",
            "комментарий",
            "примечание",
        ],

        "active": [
            "active",
            "show",
            "visible",
            "активно",
        ],
    }

    out = {}

    for key, names in (
        aliases.items()
    ):

        for name in names:

            if name in norm:

                out[key] = (
                    norm[name]
                )

                break

    return out


def val(
    row,
    idx,
    default="",
):

    if (
        idx is None
        or idx >= len(row)
    ):

        return default

    return clean_text(
        row[idx]
    )


def make_default(
    section_slug: str,
    no: int,
):

    return {

        "no":
            no,

        "id":
            f"{section_slug}-{no}",

        "title":
            "",

        "link":
            "",

        "image":
            "",

        "audio":
            "",

        "level":
            "",

        "notes":
            "",

        "active":
            True,
    }


# =========================================================
# MAIN
# =========================================================


def main():

    sections_cfg = (
        CONFIG["sections"]
    )

    cards_per = int(
        CONFIG.get(
            "cards_per_section",
            40,
        )
    )

    public_url = (
        CONFIG[
            "yandex_table_public_url"
        ]
    )

    output = {

        "generatedAt":
            datetime.now(
                timezone.utc
            ).isoformat(),

        "source":
            public_url,

        "sections":
            {},
    }

    # 1. Download fresh Yandex file.

    content = (
        download_public_file(
            public_url
        )
    )

    # 2. Repair broken XLSX table
    # references created by Yandex.

    content = (
        sanitize_broken_table_references(
            content
        )
    )

    # IMPORTANT:
    #
    # read_only=False
    #
    # is required because otherwise
    # openpyxl will NOT load pictures.

    wb = load_workbook(
        io.BytesIO(
            content
        ),
        data_only=True,
        read_only=False,
    )

    print(
        "Workbook sheets:",
        ", ".join(
            wb.sheetnames
        ),
    )

    total_populated = 0

    total_embedded = 0

    for sec in (
        sections_cfg
    ):

        slug = sec[
            "slug"
        ]

        sheet_name = sec[
            "sheet"
        ]

        cards = [

            make_default(
                slug,
                i,
            )

            for i in range(
                1,
                cards_per + 1,
            )
        ]

        if (
            sheet_name
            not in wb.sheetnames
        ):

            raise RuntimeError(
                "Required sheet "
                f"{sheet_name!r} "
                "is missing. "
                "Available sheets: "
                f"{wb.sheetnames}"
            )

        ws = wb[
            sheet_name
        ]

        rows = list(
            ws.iter_rows(
                min_row=1,
                max_row=cards_per + 1,
                values_only=True,
            )
        )

        if not rows:

            raise RuntimeError(
                f"Sheet "
                f"{sheet_name!r} "
                "is empty"
            )

        mapping = row_map(
            rows[0]
        )

        required = [
            "no",
            "title",
            "link",
        ]

        missing_headers = [

            key

            for key in required

            if key
            not in mapping
        ]

        if missing_headers:

            raise RuntimeError(
                f"Sheet "
                f"{sheet_name!r} "
                "is missing required "
                "columns: "
                f"{missing_headers}. "
                "Header row: "
                f"{rows[0]}"
            )

        # Find floating pictures.

        embedded_by_excel_row = (
            embedded_images_for_rows(
                ws,
                mapping.get(
                    "image"
                ),
                cards_per,
            )
        )

        total_embedded += len(
            embedded_by_excel_row
        )

        populated = 0

        embedded_used = 0

        for (
            position,
            row,
        ) in enumerate(
            rows[
                1:
                cards_per + 1
            ],
            start=1,
        ):

            excel_row = (
                position + 1
            )

            no_raw = val(
                row,
                mapping.get(
                    "no"
                ),
                str(position),
            )

            try:

                no = int(
                    float(
                        no_raw
                    )
                )

            except Exception:

                no = position

            if not (
                1
                <= no
                <= cards_per
            ):

                continue

            active_raw = val(
                row,
                mapping.get(
                    "active"
                ),
                "YES",
            ).lower()

            active = (
                active_raw
                not in {
                    "no",
                    "false",
                    "0",
                    "нет",
                    "off",
                }
            )

            card = {

                "no":
                    no,

                "id":
                    f"{slug}-{no}",

                "title":
                    val(
                        row,
                        mapping.get(
                            "title"
                        ),
                    ),

                "link":
                    val(
                        row,
                        mapping.get(
                            "link"
                        ),
                    ),

                "image":
                    "",

                "audio":
                    "",

                "level":
                    val(
                        row,
                        mapping.get(
                            "level"
                        ),
                    ),

                "notes":
                    val(
                        row,
                        mapping.get(
                            "notes"
                        ),
                    ),

                "active":
                    active,
            }

            # ---------------------------------
            # EMBEDDED PICTURE
            # ---------------------------------

            embedded_img = (
                embedded_by_excel_row.get(
                    excel_row
                )
            )

            if (
                embedded_img
                is not None
            ):

                try:

                    card[
                        "image"
                    ] = (
                        save_embedded_image(
                            embedded_img,
                            slug,
                            no,
                        )
                    )

                    embedded_used += 1

                except Exception as exc:

                    print(
                        "WARNING",
                        sheet_name,
                        "row",
                        excel_row,
                        ": could not extract "
                        "embedded image:",
                        exc,
                    )

            # ---------------------------------
            # IMAGE LINK FALLBACK
            # ---------------------------------

            if not card[
                "image"
            ]:

                card[
                    "image"
                ] = cache_media(
                    val(
                        row,
                        mapping.get(
                            "image"
                        ),
                    ),
                    slug,
                    no,
                    "image",
                )

            # ---------------------------------
            # AUDIO
            # ---------------------------------

            card[
                "audio"
            ] = cache_media(
                val(
                    row,
                    mapping.get(
                        "audio"
                    ),
                ),
                slug,
                no,
                "audio",
            )

            cards[
                no - 1
            ] = card

            if any(

                card.get(k)

                for k in (
                    "title",
                    "link",
                    "image",
                    "audio",
                    "notes",
                )
            ):

                populated += 1

        total_populated += (
            populated
        )

        print(
            f"{sheet_name}: "
            f"{populated} "
            "populated card(s) / "
            f"{cards_per}; "
            f"{embedded_used} "
            "embedded picture(s) used"
        )

        output[
            "sections"
        ][slug] = cards

    DATA_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    DATA_PATH.write_text(
        json.dumps(
            output,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(
        "Updated",
        DATA_PATH
    )

    print(
        "TOTAL populated cards:",
        total_populated,
    )

    print(
        "TOTAL embedded pictures mapped:",
        total_embedded,
    )

    print(
        "generatedAt:",
        output[
            "generatedAt"
        ],
    )


if __name__ == "__main__":
    main()
